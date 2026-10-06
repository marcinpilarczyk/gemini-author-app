from io import BytesIO
import re

from docx import Document

from .text import normalize_text


def manuscript_markdown(chapters) -> str:
    return "\n\n".join(
        f"## Chapter {row['chapter_num']}\n\n{row['content']}" for row in chapters
    )


def to_docx(ful_text: str, title: str) -> BytesIO:
    doc = Document()
    doc.add_heading(title, 0)

    for paragraph in normalize_text(full_text).split("\n\n"):
        if paragraph.startswith("## Chapter"):
            doc.add_heading(paragraph.replace("## ", "", 1).strip(), level=1)
            continue
        if paragraph.startswith("## "):
            doc.add_heading(paragraph.replace("## ", "", 1).strip(), level=2)
            continue

        p = doc.add_paragraph()
        parts = re.split(r"(\*\[^*]+\*\|s*[^*]+\*)", paragraph)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
